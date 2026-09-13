"""Build an idempotent, local-only import from the original three price lists."""
import hashlib
import json
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'output' / 'database-discovery'

def quoted(value):
    if value is None:
        return 'null'
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"

def prepare():
    products = {}
    sources = []
    diagnostics = []
    settings = None
    for path in sorted((ROOT / 'Archivos de excel LCP').glob('*.xlsx')):
        tier = next(t for t in ('emprendedor', 'vip', 'premium') if t in path.stem.lower())
        wb = openpyxl.load_workbook(path, data_only=False)
        assert len(wb.worksheets) == 1
        ws = wb.active
        assert [ws.cell(8, c).value for c in range(1, 4)] == ['Brand', 'Style', 'Oz']
        settings = [ws['A3'].value, ws['A4'].value, ws['A5'].value]
        source = {'filename': path.name, 'sha': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheet': ws.title, 'rows': []}
        brand = None
        for row in range(9, 269):
            raw = {openpyxl.utils.get_column_letter(c): ws.cell(row, c).value for c in range(1, 11)}
            if raw['A']:
                brand = str(raw['A']).strip()
            name = str(raw['B']).strip()
            assert brand and name and raw['G'] is None, 'Review unexpected order quantity before importing'
            size = raw['C']
            size_source = str(size)
            invalid_size = isinstance(size, datetime)
            key = hashlib.sha256(f'{brand.casefold()}|{name.casefold()}|{size_source}'.encode()).hexdigest()
            usd, nio = Decimal(str(raw['D'])), Decimal(str(raw['E']))
            assert usd > 0 and nio > 0
            if key not in products:
                match = re.search(r'HYPERLINK\("(https://[^"\s]+)"', str(raw['J']))
                products[key] = {'key': key, 'sku': f'LCP-{row-8:04d}', 'name': name, 'brand': brand,
                    'size': None if invalid_size else size, 'size_source': size_source, 'size_review': invalid_size,
                    'availability': 'sold_out' if raw['F'] == 'Agotado' else 'unspecified',
                    'image': match.group(1) if match else None, 'prices': {}}
                if invalid_size:
                    diagnostics.append({'sku': products[key]['sku'], 'name': name, 'cell': f'C{row}', 'source': size_source})
            assert tier not in products[key]['prices'], 'Duplicate product in tier'
            products[key]['prices'][tier] = {'USD': str(usd), 'NIO': str(nio)}
            source['rows'].append({'row': row, 'key': key, 'raw': raw})
        sources.append(source)
    assert len(products) == 260 and all(len(p['prices']) == 3 for p in products.values())
    sql = ['begin;']
    sql.append('insert into public.business_settings(id,name,address,phone) values(true,' + ','.join(map(quoted, settings)) + ') on conflict(id) do nothing;')
    for p in products.values():
        sql.append(f"insert into public.brands(name) values({quoted(p['brand'])}) on conflict(name) do nothing;")
        fields = [p['key'], p['sku'], p['name']]
        values = ','.join(map(quoted, fields)) + f",(select id from public.brands where name={quoted(p['brand'])})," + ','.join(map(quoted, [p['size'], 'oz', p['size_source'], p['size_review'], p['availability'], p['image']]))
        sql.append('insert into public.products(import_key,sku,name,brand_id,size,unit,size_source,size_needs_review,catalog_availability,image_reference) values(' + values + ') on conflict(import_key) do nothing;')
        product = f"(select id from public.products where import_key={quoted(p['key'])})"
        for tier, prices in p['prices'].items():
            for currency, amount in prices.items():
                sql.append(f"insert into public.product_prices(product_id,tier_code,currency,amount) values({product},{quoted(tier)},{quoted(currency)},{amount}) on conflict do nothing;")
        for location in ('store', 'warehouse'):
            sql.append(f"insert into public.inventory_balances(product_id,location,quantity) values({product},{quoted(location)},null) on conflict do nothing;")
    for s in sources:
        sql.append('insert into private.import_sources(filename,sha256,sheet_name,row_count) values(' + ','.join(map(quoted, [s['filename'], s['sha'], s['sheet'], len(s['rows'])])) + ') on conflict(sha256) do nothing;')
        for r in s['rows']:
            sql.append(f"insert into private.import_rows(source_id,row_number,product_id,raw_values) values((select id from private.import_sources where sha256={quoted(s['sha'])}),{r['row']},(select id from public.products where import_key={quoted(r['key'])}),{quoted(json.dumps(r['raw'],ensure_ascii=False,default=str))}::jsonb) on conflict do nothing;")
    sql.append('commit;')
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / 'catalog-import.sql').write_text('\n'.join(sql), encoding='utf-8')
    chunks, current = [], []
    for statement in sql[1:-1]:
        if sum(map(len,current)) + len(statement) > 32000:
            chunks.append(current)
            current = []
        current.append(statement)
    if current:
        chunks.append(current)
    for i, statements in enumerate(chunks):
        (OUTPUT / f'import-{i:02d}.sql').write_text('begin;\n'+'\n'.join(statements)+'\ncommit;',encoding='utf-8')
    print('Import chunks:',len(chunks))
    (OUTPUT / 'catalog.json').write_text(json.dumps(list(products.values()), ensure_ascii=False, indent=2), encoding='utf-8')
    report = {'products': len(products), 'brands': len({p['brand'] for p in products.values()}), 'prices': 1560,
              'source_rows': sum(len(s['rows']) for s in sources), 'size_review': diagnostics,
              'stock': 'not provided', 'sales': 'not provided', 'customers': 'not provided'}
    (OUTPUT / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))

if __name__ == '__main__':
    prepare()
